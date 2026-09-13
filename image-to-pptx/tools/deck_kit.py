# -*- coding: utf-8 -*-
"""按源图像素坐标搭 pptx 的工具层。

用法：
    from deck_kit import *
    canvas(1673, 951)          # 按源图像素设坐标系，幻灯片固定 13.333x7.5in
    prs = new_deck()
    sh  = add_slide(prs)
    ...
    save(prs, "out.pptx")      # 内含主题字体统一

坐标一律写源图像素，X()/Y() 负责换算到 EMU——脚本里的数字可以直接和图对照。
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree

FONT = "微软雅黑"
SLIDE_W_IN, SLIDE_H_IN = 13.3333, 7.5
_C = {"w": 1673.0, "h": 951.0}

def canvas(px_w, px_h):
    """按源图尺寸设坐标系。幻灯片始终 16:9，源图比例不同则纵向轻微压缩（1% 内不可察）。"""
    _C["w"], _C["h"] = float(px_w), float(px_h)

def set_font(name):
    global FONT
    FONT = name

def X(px):  return Inches(px * SLIDE_W_IN / _C["w"])
def Y(px):  return Inches(px * SLIDE_H_IN / _C["h"])
W = X
H = Y
def pt_of(px):  return px * 960.0 / _C["w"]          # 源图像素 -> 磅
def wpx(text, pt):
    """粗估文本像素宽：CJK 记 1 字宽，拉丁记 0.5。

    回落字体（本机无微软雅黑时）比雅黑宽 7~12%，排版留余量按这个系数。
    """
    return sum(0.5 if ord(c) < 0x2E80 else 1.0 for c in text) * pt / (960.0 / _C["w"])

# ---------- 中性配色，各 deck 按源图覆盖 ----------
INK   = RGBColor(0x1A,0x1A,0x1A); BODY  = RGBColor(0x33,0x3A,0x45)
GRAY  = RGBColor(0x8A,0x93,0xA0); WHITE = RGBColor(0xFF,0xFF,0xFF)
SEP   = RGBColor(0xD9,0xDF,0xE6)
BLUE  = RGBColor(0x2C,0x72,0xC4); BLUE_L  = RGBColor(0xE9,0xF1,0xFB); BLUE_B  = RGBColor(0xAF,0xCC,0xEA)
BLUE_D= RGBColor(0x15,0x65,0xC0)
GREEN = RGBColor(0x4C,0x9E,0x3F); GREEN_L = RGBColor(0xEC,0xF6,0xE7); GREEN_B = RGBColor(0xB6,0xDB,0xAB)
ORANGE= RGBColor(0xEF,0x7D,0x1A); ORANGE_L= RGBColor(0xFD,0xF2,0xE8); ORANGE_B= RGBColor(0xF6,0xC9,0x9B)
RED   = RGBColor(0xE0,0x52,0x52); RED_L   = RGBColor(0xFD,0xF0,0xF0); RED_B   = RGBColor(0xF2,0xB8,0xB8)
PURPLE= RGBColor(0x6B,0x3F,0xA0); PURPLE_L= RGBColor(0xF3,0xEF,0xFB); PURPLE_B= RGBColor(0xC2,0xAC,0xE3)

def mix(c1, c2, t):
    """c1 向 c2 混色。RGBColor 是 tuple 子类，可直接下标。"""
    return RGBColor(*(int(round(c1[i] + (c2[i] - c1[i]) * t)) for i in range(3)))

# ---------- 字体 + 关拼写检查 ----------
def style_run(run, size, bold=False, color=BODY, italic=False):
    """四个字体槽位全写 + noProof。两件独立的事，别混为一谈：

    - **红色波浪线**来自拼写检查，跟字体无关。`noProof="1"` 才是开关；
      只设 `lang="zh-CN"` 不保险（编辑语言仍可能触发校对）。
    - **中文字形回落**来自字体槽位：`font.name` 只写 `a:latin`，中文走 `a:ea`，
      漏了就用主题字体渲染。四槽位 latin/ea/cs/sym 全写才锁得住。
    """
    f = run.font
    f.size = Pt(size); f.bold = bold; f.italic = italic
    f.color.rgb = color
    f.name = FONT                       # 写 a:latin，位置由 python-pptx 保证
    rPr = run._r.get_or_add_rPr()
    rPr.set('lang', 'zh-CN'); rPr.set('altLang', 'en-US')
    rPr.set('noProof', '1'); rPr.set('dirty', '0')
    latin = rPr.find(qn('a:latin')); idx = list(rPr).index(latin)
    for i, tag in enumerate(('a:ea', 'a:cs', 'a:sym')):
        for e in rPr.findall(qn(tag)): rPr.remove(e)
        el = etree.SubElement(rPr, qn(tag)); rPr.remove(el)
        el.set('typeface', FONT); rPr.insert(idx + 1 + i, el)
    return run

def noshadow(sh):
    try: sh.shadow.inherit = False
    except Exception: pass
    return sh

# ---------- 形状 ----------
def _paint(s, fill, line_c, lw):
    noshadow(s)
    if fill is None: s.fill.background()
    else: s.fill.solid(); s.fill.fore_color.rgb = fill
    if line_c is None: s.line.fill.background()
    else: s.line.color.rgb = line_c; s.line.width = Pt(lw)
    s.text_frame.word_wrap = False
    return s

def rrect(sh, x, y, w, h, fill=None, line=None, lw=1.0, radius=8):
    s = sh.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, X(x), Y(y), W(w), H(h))
    try: s.adjustments[0] = max(0.0, min(0.5, radius / float(min(w, h))))
    except Exception: pass
    return _paint(s, fill, line, lw)

def rect(sh, x, y, w, h, fill=None, line=None, lw=1.0):
    return _paint(sh.add_shape(MSO_SHAPE.RECTANGLE, X(x), Y(y), W(w), H(h)), fill, line, lw)

def oval(sh, x, y, w, h, fill=None, line=None, lw=1.0):
    return _paint(sh.add_shape(MSO_SHAPE.OVAL, X(x), Y(y), W(w), H(h)), fill, line, lw)

def shape(sh, kind, x, y, w, h, fill=None, line=None, lw=1.0, rot=None):
    s = sh.add_shape(kind, X(x), Y(y), W(w), H(h))
    if rot is not None: s.rotation = rot
    return _paint(s, fill, line, lw)

def _lnstyle(c, color, lw, arrow, back, dash):
    c.line.color.rgb = color; c.line.width = Pt(lw)
    ln = c.line._get_or_add_ln()
    if dash:
        d = etree.SubElement(ln, qn('a:prstDash')); d.set('val', dash)
    if back:
        e = etree.SubElement(ln, qn('a:headEnd')); e.set('type','triangle'); e.set('w','med'); e.set('len','med')
    if arrow:
        e = etree.SubElement(ln, qn('a:tailEnd')); e.set('type','triangle'); e.set('w','med'); e.set('len','med')
    return c

def line(sh, x1, y1, x2, y2, color=SEP, lw=1.0, arrow=False, back=False, dash=None):
    """直线连接符。箭头用原生 line-end——直线走 prstGeom，各渲染器都认。"""
    c = sh.add_connector(MSO_CONNECTOR.STRAIGHT, X(x1), Y(y1), X(x2), Y(y2))
    return _lnstyle(c, color, lw, arrow, back, dash)

def poly(sh, x, y, w, h, pts, fill=None, lncolor=None, lw=1.0, close=True):
    """自由多边形。pts 为 0..1 归一化坐标，按 (x,y,w,h) 包围盒展开。"""
    ax = [(x + px * w, y + py * h) for px, py in pts]
    ff = sh.build_freeform(X(ax[0][0]), Y(ax[0][1]))
    ff.add_line_segments([(X(px), Y(py)) for px, py in ax[1:]], close=close)
    return _paint(ff.convert_to_shape(), fill, lncolor, lw)

def bez(sh, x1, y1, x2, y2, color, lw=1.4, arrow=True, bow=0.5, n=30):
    """水平出入的贝塞尔曲线箭头。

    箭头画成实心三角而不是挂 a:tailEnd：QuickLook / 预览 / Keynote 不给
    custGeom 开放路径渲染 line-end，挂上去箭头会整个消失（PowerPoint 里正常）。
    """
    AH = 10.0 if arrow else 0.0
    ex = x2 - AH
    c1x, c1y = x1 + (ex - x1) * bow, y1
    c2x, c2y = ex - (ex - x1) * bow, y2
    pts = []
    for i in range(n + 1):
        t = i / float(n); m = 1 - t
        pts.append((m**3*x1 + 3*m*m*t*c1x + 3*m*t*t*c2x + t**3*ex,
                    m**3*y1 + 3*m*m*t*c1y + 3*m*t*t*c2y + t**3*y2))
    ff = sh.build_freeform(X(pts[0][0]), Y(pts[0][1]))
    ff.add_line_segments([(X(a), Y(b)) for a, b in pts[1:]], close=False)
    s = ff.convert_to_shape()
    noshadow(s); s.fill.background()
    s.line.color.rgb = color; s.line.width = Pt(lw)
    s.text_frame.word_wrap = False
    if arrow:
        tri = sh.build_freeform(X(x2), Y(y2))
        tri.add_line_segments([(X(x2 - AH), Y(y2 - 4.6)), (X(x2 - AH), Y(y2 + 4.6))], close=True)
        _paint(tri.convert_to_shape(), color, None, 0)
    return s

def dashbox(sh, x, y, w, h, color, fill=None, radius=6, lw=1.1):
    s = rrect(sh, x, y, w, h, fill, color, lw, radius)
    d = etree.SubElement(s.line._get_or_add_ln(), qn('a:prstDash')); d.set('val', 'dash')
    return s

# ---------- 文本 ----------
def tb(sh, x, y, w, h, lines, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE, wrap=True, spacing=None):
    """多行文本框。lines = [(text, size, bold, color)]，每行一个段落。"""
    box = sh.add_textbox(X(x), Y(y), W(w), H(h))
    tf = box.text_frame
    tf.word_wrap = wrap
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    for i, item in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if spacing: p.line_spacing = spacing
        r = p.add_run(); r.text = item[0]
        style_run(r, item[1], item[2], item[3])
    return box

def label(sh, x, y, w, h, text, size, bold=False, color=BODY,
          align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE):
    return tb(sh, x, y, w, h, [(text, size, bold, color)], align, anchor, wrap=False)

def rich(sh, x, y, w, h, parts, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE, wrap=False):
    """同一段落内混排多个 run：parts = [(text, size, bold, color)]。"""
    box = sh.add_textbox(X(x), Y(y), W(w), H(h))
    tf = box.text_frame
    tf.word_wrap = wrap; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]; p.alignment = align
    for t, size, bold, color in parts:
        r = p.add_run(); r.text = t; style_run(r, size, bold, color)
    return box

# ---------- deck ----------
def new_deck():
    prs = Presentation()
    prs.slide_width  = Inches(SLIDE_W_IN)
    prs.slide_height = Inches(SLIDE_H_IN)
    return prs

def add_slide(prs, bg=WHITE):
    slide = prs.slides.add_slide(prs.slide_layouts[6])   # 空白版式
    sh = slide.shapes
    if bg is not None:
        rect(sh, 0, 0, _C["w"], _C["h"], bg, None)
    return sh

def save(prs, path):
    """存盘前把主题的 majorFont / minorFont 也统一，兜住未经 style_run 的文本。"""
    theme = prs.slide_masters[0].part.part_related_by(
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/theme")
    t = etree.fromstring(theme.blob)
    ns = {'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'}
    for scheme in ('majorFont', 'minorFont'):
        node = t.find('.//a:fontScheme/a:%s' % scheme, ns)
        for tag in ('latin', 'ea', 'cs'):
            e = node.find('a:%s' % tag, ns)
            if e is not None: e.set('typeface', FONT)
    theme._blob = etree.tostring(t, xml_declaration=True, encoding='UTF-8', standalone=True)
    prs.save(path)
    return path
